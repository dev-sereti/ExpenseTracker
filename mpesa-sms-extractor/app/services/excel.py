from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Transaction

HEADERS = ["Transaction Code", "Date", "Amount", "Mode", "Sender", "Receiver", "Raw Message"]


def transactions_to_xlsx(transactions: list[Transaction]) -> bytes:
    wb = Workbook()
    active = wb.active
    if active is None or not isinstance(active, Worksheet):
        ws: Worksheet = wb.create_sheet()
    else:
        ws = active

    ws.title = "Transactions"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True)

    for tx in transactions:
        ws.append(
            [
                tx.transaction_code,
                tx.transaction_date.isoformat(sep=" ", timespec="minutes"),
                float(tx.amount),
                tx.transaction_mode,
                tx.sender_name or "",
                tx.receiver_name or "",
                tx.raw_message,
            ]
        )

    widths = [18, 20, 12, 12, 22, 22, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def dashboard_report_to_xlsx(
    *,
    months: list[str],
    monthly: dict[str, list[float]],
    cumulative: dict[str, list[float]],
    totals: dict[str, float],
    window_start: str,
    window_end: str,
    transactions: list[Transaction],
) -> bytes:
    """
    Creates an .xlsx "dashboard report" for the last 3 months:
    - Summary sheet: totals + monthly + cumulative tables
    - Transactions sheet: raw transactions in the same window
    """
    wb = Workbook()

    active = wb.active
    if active is None or not isinstance(active, Worksheet):
        ws_summary: Worksheet = wb.create_sheet()
    else:
        ws_summary = active
    ws_summary.title = "Dashboard Summary"

    # Styles
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="243055")  # matches your theme border-ish
    header_align = Alignment(horizontal="left", vertical="center")

    def write_header_row(ws: Worksheet, row: int, values: list[str]) -> None:
        ws.append(values)
        for col_idx in range(1, len(values) + 1):
            c = ws.cell(row=row, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align

    # Title / window info
    ws_summary["A1"] = "M-PESA Dashboard Report"
    ws_summary["A1"].font = title_font

    ws_summary["A3"] = "Period"
    ws_summary["B3"] = f"{window_start} → {window_end}"

    # Totals section
    ws_summary["A5"] = "Totals (last 3 months)"
    ws_summary["A5"].font = Font(bold=True)

    write_header_row(ws_summary, 6, ["Metric", "Amount (Ksh)"])
    totals_rows = [
        ("Received", totals.get("Received", 0.0)),
        ("Sent", totals.get("Sent", 0.0)),
        ("Paid", totals.get("Paid", 0.0)),
        ("Net (Received - Sent - Paid)", totals.get("Net", 0.0)),
    ]
    for metric, value in totals_rows:
        ws_summary.append([metric, value])

    # Format amounts in column B (totals)
    for r in range(7, 7 + len(totals_rows)):
        ws_summary[f"B{r}"].number_format = '#,##0.00'

    # Monthly section
    ws_summary["A11"] = "Monthly totals"
    ws_summary["A11"].font = Font(bold=True)

    write_header_row(ws_summary, 12, ["Month", "Received", "Sent", "Paid", "Net"])
    for i, month in enumerate(months):
        ws_summary.append(
            [
                month,
                monthly.get("Received", [0, 0, 0])[i],
                monthly.get("Sent", [0, 0, 0])[i],
                monthly.get("Paid", [0, 0, 0])[i],
                monthly.get("Net", [0, 0, 0])[i],
            ]
        )

    # Numeric format for monthly section
    for r in range(13, 16):
        for col in ["B", "C", "D", "E"]:
            ws_summary[f"{col}{r}"].number_format = '#,##0.00'

    # Cumulative section
    ws_summary["A17"] = "Cumulative totals (running across the 3 months)"
    ws_summary["A17"].font = Font(bold=True)

    write_header_row(ws_summary, 18, ["Month", "Cum Received", "Cum Sent", "Cum Paid", "Cum Net"])
    for i, month in enumerate(months):
        ws_summary.append(
            [
                month,
                cumulative.get("Received", [0, 0, 0])[i],
                cumulative.get("Sent", [0, 0, 0])[i],
                cumulative.get("Paid", [0, 0, 0])[i],
                cumulative.get("Net", [0, 0, 0])[i],
            ]
        )

    for r in range(19, 22):
        for col in ["B", "C", "D", "E"]:
            ws_summary[f"{col}{r}"].number_format = '#,##0.00'

    # Column widths for summary
    ws_summary.column_dimensions["A"].width = 34
    ws_summary.column_dimensions["B"].width = 22
    ws_summary.column_dimensions["C"].width = 16
    ws_summary.column_dimensions["D"].width = 16
    ws_summary.column_dimensions["E"].width = 16

    # Freeze panes so headers stay visible
    ws_summary.freeze_panes = "A7"

    # Transactions sheet (same window)
    ws_tx: Worksheet = wb.create_sheet("Transactions")
    ws_tx.append(HEADERS)
    for cell in ws_tx[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)

    for tx in transactions:
        ws_tx.append(
            [
                tx.transaction_code,
                tx.transaction_date.isoformat(sep=" ", timespec="minutes"),
                float(tx.amount),
                tx.transaction_mode,
                tx.sender_name or "",
                tx.receiver_name or "",
                tx.raw_message,
            ]
        )

    widths = [18, 20, 12, 12, 22, 22, 60]
    for i, w in enumerate(widths, start=1):
        ws_tx.column_dimensions[get_column_letter(i)].width = w
    ws_tx.freeze_panes = "A2"

    # Amount formatting in tx sheet (column C)
    for r in range(2, 2 + len(transactions)):
        ws_tx[f"C{r}"].number_format = '#,##0.00'

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()